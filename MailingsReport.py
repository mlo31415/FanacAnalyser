from __future__ import annotations

# ======================================================================================================================
# The APA mailings reports.
#
# This was the separate FanacMailings program until 2026-09-17.  It used to run afterwards, reading back the
# mailings.csv that FanacAnalyser writes; it now runs as FanacAnalyser's last report, working directly from the
# FanzineIssueInfo list that has already been built.  (mailings.csv is still written -- it is a useful artifact --
# but nothing reads it any more.)
#
# It combines two sources:
#   * the mailings FanacAnalyser found on the fanzine index pages (which apazine was in which mailing)
#   * Joe's "APA Mailings.xlsx" (who the OE was and when each mailing came out), which FanacAnalyser has no other use for
# and writes, under <reports>/APAs/:
#   * one page per mailing          <apa>/<number>.html
#   * one page per APA              <apa>/index.html
#   * one page listing all the APAs  index.html
# ======================================================================================================================

from dataclasses import dataclass, field
import os
import re
import datetime

import openpyxl

from FanzineIssueSpecPackage import FanzineDate, FanzineIssueInfo
from Settings import Settings
from HelpersPackage import FindAndReplaceBracketedText, ParseFirstStringBracketedText, SortMessyNumber, SortTitle, Pluralize, NormalizePersonsName, Int0, FormatLink
from HelpersPackage import FindIndexOfStringInList, FormatCount, UnicodeToHtml, MakeFancyLink, SplitOnAnySingleChar, RemoveNonAlphanumericChars
from Log import LogError, Log


# =============================================================================
# Reduce an APA's name to a form which compares equal regardless of spaces, punctuation or case, so that
# "Shadow FAPA", "Shadow-FAPA" and "ShadowFAPA" are all recognized as the same APA.
def CompressAPAName(name: str) -> str:
    return RemoveNonAlphanumericChars(name.casefold()).replace(" ", "")


# =============================================================================
# The APA's name is also its directory name and part of the link to it, but Windows forbids \ / : * ? " < > | in a
# file name and "APA:NESFA" contains one.  Those characters become "-"; every other APA name is returned unchanged.
# The name itself is still displayed as it is spelled in the setting -- this is only for paths and hrefs.
def APADirName(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "-", name)


# =============================================================================
# The full path of the parameters file the settings were loaded from, for use in error messages.
# Settings().Dictpath is a property which is "" until a settings file has been loaded, so it is always safe to read.
def SettingsFileName() -> str:
    return Settings().Dictpath or "the parameters file"


# =============================================================================
# Generate all of the APA mailing pages.
# fanacIssueList is the complete list of issues FanacAnalyser has already read.
# rootDir is where the templates, the bumpf files and the xlsx live; apaReportsDir is where the pages are written.
def GenerateMailingsReports(fanacIssueList: list[FanzineIssueInfo], rootDir: str, apaReportsDir: str) -> None:

    # **************************************************************************
    # Get the list of known apas.  Settings().Get() returns None -- not "" -- when the setting is missing altogether.
    knownApas=Settings().Get("Known APAs", "")
    if not knownApas:
        LogError(f"***The APA mailings report was skipped: no 'Known APAs' setting (the list of APAs to report on)"
                 f" was found in {SettingsFileName()}")
        return
    knownApas=[x.replace('"', '').strip() for x in knownApas.split(",")]

    # APA names are matched with spaces, punctuation and case ignored, so that a page writing "ShadowFAPA 12" or
    # "Shadow-FAPA 12" is recognized as the "Shadow FAPA" of the setting -- those are spellings of one APA, not two.
    # Two entries of the setting which compress to the same thing are therefore almost certainly one APA listed twice.
    # We cannot tell which was meant, so neither is matched loosely and the duplicate is reported for removal.
    apasByCompressedName: dict[str, str]={}
    ambiguous: set[str]=set()
    for apaName in knownApas:
        key=CompressAPAName(apaName)
        if key in apasByCompressedName and apasByCompressedName[key] != apaName:
            LogError(f"***APA mailings: the Known APAs setting in {SettingsFileName()} lists both"
                     f" '{apasByCompressedName[key]}' and '{apaName}'.  Those are the same once spaces, punctuation and case"
                     f" are ignored, so they are two spellings of one APA and one of them should be deleted from the setting."
                     f"  Until then neither is matched loosely, and a mailing has to spell one of them exactly to be counted.")
            ambiguous.add(key)
        apasByCompressedName[key]=apaName
    for key in ambiguous:
        del apasByCompressedName[key]

    # **************************************************************************
    # for each known apa, read Joe's APA mailings data if it exists
    # Mailings is a dictionary indexed by the apa name.
    #   The value is a dictionary indexed by the mailing number as a string
    #       The value of *that* is a MailingDev
    # Note that we do  not fill in Counts here
    mailingsInfoTablefromJoe: dict[str, dict[str, MailingInfoFromJoe]]={}
        # 1st level key is APA name
        # 2nd level key is mailing name
    # Check for the spreadsheet once rather than once per APA, so a missing one is reported once and not two dozen times
    xlsxPath=os.path.abspath(os.path.join(rootDir, "APA Mailings.xlsx"))
    if not os.path.exists(xlsxPath):
        LogError(f"***APA mailings: {xlsxPath} (Joe's table of mailing dates and Official Editors) was not found."
                 f"  The mailing pages will be generated without dates or OEs.")
    for apaName in knownApas:
        table=ReadXLSX(xlsxPath, apaName) if os.path.exists(xlsxPath) else None
        if table is None:
            table={}
        mailingsInfoTablefromJoe[apaName]=table

    # ---------------------------
    # Sort the issues into a dictionary of the form dict(apa, dict(mailing, data))
    # Allmailings is keyed by the apa's name.  The value is an EntireAPA object
    allAPAs: AllAPAs=AllAPAs()
    numIssues=0
    for fanzine in fanacIssueList:
        # Select only issues which have an entry in the mailings column
        if len(fanzine.Mailings) == 0:
            continue
        numIssues+=1
        for entry in fanzine.Mailings:
            # A single entry may name more than one mailing: 'FAPA 20 & VAPA 23'
            for mailing in [x.strip() for x in SplitOnAnySingleChar("&,", entry)]:
                # A mailing is "<apa name> <number>".  The name itself may contain a space ("Shadow FAPA 6") and so may
                # the number ("FAPA 110 postmailing"), so try the longest name which matches: otherwise "Shadow FAPA 6"
                # would be read as the APA "Shadow", and "FAPA 110 postmailing" as a mailing of some APA "FAPA 110".
                tokens=mailing.split()
                for n in range(len(tokens)-1, 0, -1):
                    name=" ".join(tokens[:n])
                    apaName=apasByCompressedName.get(CompressAPAName(name))
                    if apaName is None and name in knownApas:   # An ambiguous name has to be spelled exactly
                        apaName=name
                    if apaName is not None:
                        allAPAs[apaName][" ".join(tokens[n:])].append(fanzine)
                        break

    if numIssues < 100:
        LogError(f"There are {numIssues} issues with mailing information -- there should be many hundreds")

    # ------------------
    # We've slurped in all the data.
    # Now merge Joe's mailing info into allAPAs
    for apa in allAPAs:
        for mailing in apa:
            if apa.Name in mailingsInfoTablefromJoe:
                if mailing.Number in mailingsInfoTablefromJoe[apa.Name]:
                    mailing.MIFJ=mailingsInfoTablefromJoe[apa.Name][mailing.Number]


    # The next step is to generate the counts
    # Walk through allAPAs
    # For each APA that we found there extract the data, merge it was Joe's data, and create a unified dataset to generate the web pages
    countAllAPAs=Counts()       # This is the only 'bare' Counts -- all the others are in larger structures
    for apa in allAPAs:

        # For each mailing of that APA count up the issues and pages
        for mailing in apa:
            # Now generate the data rows in the mailings table
            for apazine in mailing:
                mailing.Count+=Counts(Issues=1, Pages=apazine.Pagecount)
            apa.Count+=Counts(Mailings=1, Issues=mailing.Count.Issues, Pages=mailing.Count.Pages)

        countAllAPAs+=Counts(Mailings=apa.Count.Mailings, Issues=apa.Count.Issues, Pages=apa.Count.Pages)


    ##################################################################################################################
    # We have done all the analysis: generate the HTML pages

    # We will create a directory in the APA reports dir for each APA, and put the individual issue index pages there
    if not os.path.exists(apaReportsDir):
        os.makedirs(apaReportsDir)

    # All the pages we generate here need the same kinds of information to be added:
    #   Page title
    #   Page metadata
    #   Updated timestamp
    def AddBoilerplate(page: str, title: str, metadata: str) -> str:
        start, mid, end=ParseFirstStringBracketedText(page, "fanac-title")
        mid=mid.replace("title of page", title)
        page=start+mid+end

        start, mid, end=ParseFirstStringBracketedText(page, "head")
        mid=mid.replace("mailing content", metadata)
        page=start+mid+end

        # Add the updated date/time
        page, _=FindAndReplaceBracketedText(page, "fanac-updated", f"Updated {datetime.datetime.now().strftime('%m/%d/%Y, %H:%M:%S')}")
        return page

    # Read one of the three template files.  They live alongside the other control files in rootDir.
    def ReadTemplate(settingName: str, description: str) -> str|None:
        fname=Settings().Get(settingName, "")
        if not fname:
            LogError(f"***The APA mailings report was skipped: no '{settingName}' setting (naming the {description})"
                     f" was found in {SettingsFileName()}")
            return None
        path=os.path.abspath(os.path.join(rootDir, fname))
        try:
            with open(path, "r", encoding="utf-8") as file:
                return "".join(file.readlines())
        except FileNotFoundError:
            LogError(f"***The APA mailings report was skipped: the {description} was not found."
                     f"  '{settingName}' in {SettingsFileName()} names '{fname}', so I looked for {path}")
            return None

    templateMailing=ReadTemplate("Template-Mailing", "template file for an individual mailing page")
    if templateMailing is None:
        return
    templateApa=ReadTemplate("Template-APA", "template for an APA page")
    if templateApa is None:
        return

    # Walk through the info generated by FanacAnalyser.
    # For each APA that we found there:
    #   Create an apa HTML page listing (and linking to) all the mailing pages
    #   Create all the individual mailing pages
    allAPAs.sort()
    for apa in allAPAs:

        # Make sure that a directory exists for this APA's html files
        if not os.path.exists(os.path.join(apaReportsDir, APADirName(apa.Name))):
            os.mkdir(os.path.join(apaReportsDir, APADirName(apa.Name)))

        apa.sort()
        for mailing in apa:
            mailing.sort()

            ##################################################################
            ##################################################################
            # Do a mailing page
            # First, the top matter
            # <div><fanac-top>
            # <table class=topmatter>
            # <tr><td class=topmatter>mailing</td></tr>
            # <tr><td class=topmatter>editor</td></tr>
            # <tr><td class=topmatter>date</td></tr>
            # </table>
            # </fanac-top></div>
            mailingPage=templateMailing
            start, mid, end=ParseFirstStringBracketedText(mailingPage, "fanac-top")
            editor=f"OE: {NormalizePersonsName(mailing.MIFJ.Editor)}"
            when=mailing.MIFJ.Date.FormatDate("%B %Y")
            number=mailing.Number
            mid=mid.replace("editor", editor)
            mid=mid.replace("date", when)
            mid=mid.replace("mailing", f"{apa.Name} Mailing #{number}")
            mailingPage=start+mid+end

            mailingPage=AddBoilerplate(mailingPage, f"{apa.Name}-{mailing.Number}", f"{mailing.Number}, {editor}, {when}, {apa.Name}-mailing")

            # Now the bottom matter (the list of fanzines)
            newtable="<tr>\n"
            newtable+="<th>Contribution</th>\n"
            newtable+="<th>Editor</th>\n"
            newtable+="<th>Pages</th>\n"
            newtable+="</tr>\n"

            # Now generate the data rows in the mailings table
            for apazine in mailing:
                newtable+="<tr>\n"
                if apazine.DirURL != "" and apazine.PageFilename != "":
                    if apazine.PageFilename.startswith("//fanac.org"):
                        # It's an absolute reference
                        href=apazine.PageFilename
                    else:
                        # It's a relative reference
                        href=f"{apazine.DirURL}/{apazine.PageFilename}"
                    href=href.replace(" ", "%20")
                    newtable+=f"<td>{FormatLink(href, UnicodeToHtml(apazine.IssueName))}</td>\n"
                else:
                    newtable+=f"<td>&nbsp;</td>\n"
                if apazine.Editor != "":
                    newtable+=f"<td>{MakeFancyLink(apazine.Editor)}&nbsp;&nbsp;</td>"
                else:
                    newtable+=f"<td>&nbsp;</td>\n"
                if apazine.Pagecount:   # Pagecount is an int now, and 0 means "not known"
                    newtable+=f"<td>{apazine.Pagecount}</td>\n"
                else:
                    newtable+=f"<td>&nbsp;</td>\n"
                newtable+="</tr>\n"
            newtable=newtable.replace("\\", "/")

            # Insert the new issues table into the template
            mailingPage, success=FindAndReplaceBracketedText(mailingPage, "fanac-rows", newtable)
            if not success:
                LogError(f"Could not add issues table to the mailing page at 'fanac-rows'")
                return

            # Insert the label for the button taking you to the previous mailing for this APA
            index=apa.prevIndex(mailing.Number)
            if index is None:
                buttonText=f"No prev mailing "
                link=""
            else:
                buttonText=f" Prev Mailing (#{apa[index].Number}) "
                link=f'"{apa[index].Number}.html"'
            mailingPage, success=FindAndReplaceBracketedText(mailingPage, "fanac-PrevMailing", buttonText)
            if success:
                mailingPage=mailingPage.replace('"prev.html"', link)
            if not success:
                LogError(f"Could not change prev button text on the mailing page at 'fanac-PrevMailing'")
                return

            # Insert the label for the button taking you up one level to all mailings for this APA
            mailingPage, success=FindAndReplaceBracketedText(mailingPage, "fanac-AllMailings", f"All {apa.Name} mailings")
            if not success:
                LogError(f"Could not change up to APA button text on the mailing page at 'fanac-AllMailings'")
                return

            # Insert the label for the button taking you to the next mailing for this APA
            index=apa.nextIndex(mailing.Number)
            if index is None:
                buttonText=f"No next mailing "
                link=""
            else:
                buttonText=f" Next Mailing (#{apa[index].Number}) "
                link=f'"{apa[index].Number}.html"'
            mailingPage, success=FindAndReplaceBracketedText(mailingPage, "fanac-NextMailing", buttonText)
            if success:
                mailingPage=mailingPage.replace('"next.html"', link)
            if not success:
                LogError(f"Could not change next button text on the mailing page at 'fanac-NextMailing'")
                return

            # Modify the Mailto: so that the page name appears as the subject
            mailingPage, success=FindAndReplaceBracketedText(mailingPage, "fanac-ThisPageName", f"{apa.Name}:{mailing.Number}")
            if not success:
                LogError(f"Could not change mailto Subject on the mailing page at 'fanac-ThisPageName'")
                #return

            # Add counts of mailings and contributions to bottom
            start, mid, end=ParseFirstStringBracketedText(mailingPage, "fanac-totals")
            mailingPage=f"{start} {mailing.Count}  {end}"

            # Write the mailing file
            fn=os.path.join(apaReportsDir, APADirName(apa.Name), mailing.Number)+".html"
            with open(fn, "w", encoding="utf-8") as file:
                mailingPage=mailingPage.split("\n")
                file.writelines(mailingPage)


        ##################################################################
        ##################################################################
        # Now that the mailing pages are all done, do an apa page

        # Add the APA's name at the top
        start, mid, end=ParseFirstStringBracketedText(templateApa, "fanac-top")
        mid=mid.replace("apa-name", apa.Name)
        newAPAPage=start+mid+end

        # Add random descriptive information if a file <apa>-bumpf.txt exists.  (E.g., SAPS-bumpf.txt)
        fname=os.path.join(rootDir, APADirName(apa.Name)+"-bumpf.txt")
        if os.path.exists(fname):
            with open(fname, "r", encoding="utf-8") as file:
                bumpf=file.read()
            if len(bumpf) > 0:
                start, mid, end=ParseFirstStringBracketedText(newAPAPage, "fanac-bumpf")
                if len(end) > 0:
                    mid=bumpf+"<p>"
                    newAPAPage=start+mid+end
            Log(f"Bumpf added to {apa.Name} page")
        else:
            Log(f" No {apa.Name}-bumpf.txt file found, so no bumpf added to {apa.Name} page.")

        newAPAPage=AddBoilerplate(newAPAPage, f"{apa.Name} Mailings", f"{apa.Name} mailings")

        loc=newAPAPage.find("</fanac-rows>")
        if loc < 0:
            LogError(f"The APA template is missing the '</fanac-rows>' indicator.")
            return
        newAPAPageFront=newAPAPage[:loc]
        newAPAPageRear=newAPAPage[loc+len("</fanac-rows>"):]

        apa.sort()
        for mailing in apa:
            when=mailing.MIFJ.Date
            editor=mailing.MIFJ.Editor
            issues=mailing.Count.Issues
            pages=mailing.Count.Pages
            newAPAPageFront+=(f"\n<tr><td>{FormatLink(mailing.Number+".html", mailing.Number)}</td>"
                              f"<td>{when}</td><td>{editor}</td>"
                              f"<td style='text-align: right'>{issues}&nbsp;&nbsp;&nbsp;&nbsp;</td>"
                              f"<td style='text-align: right'>{pages}&nbsp;&nbsp;&nbsp;&nbsp;</td>"
                              f"</tr>")

        newAPAPage=newAPAPageFront+newAPAPageRear

        # Add counts of mailings and contributions to bottom
        start, mid, end=ParseFirstStringBracketedText(newAPAPage, "fanac-totals")
        newAPAPage=f"{start} {apa.Count}  {end}"

        # Add the updated date/time
        newAPAPage, success=FindAndReplaceBracketedText(newAPAPage, "fanac-updated", f"Updated {datetime.datetime.now().strftime('%m/%d/%Y, %H:%M:%S')}>")

        # Make the mailto correctly list the apa in the subject line
        newAPAPage, success=FindAndReplaceBracketedText(newAPAPage, "fanac-APAPageMailto", f"Issue related to APA {apa.Name}")
        if not success:
            LogError(f"The APA template is missing the '</fanac-APAPageMailto>' indicator.")
            return

        # Write out the APA list of all mailings
        with open(os.path.join(apaReportsDir, APADirName(apa.Name), "index.html"), "w", encoding="utf-8") as file:
            file.writelines(newAPAPage)

    ##################################################################
    ##################################################################
    # Generate the All Apas root page

    templateAllApas=ReadTemplate("Template-allAPAs", "template for the page listing all APAs")
    if templateAllApas is None:
        return

    templateAllApas=AddBoilerplate(templateAllApas, f"Mailings for All APAs", f"Mailings for All APAs")

    listText="\n<i>Click on the APA's name to see APA's contents</i>\n"
    listText+="<style>th, td{border-style: hidden;}</style>\n\n"

    listText+="<table>\n<tr>\n<th>&nbsp;&nbsp;&nbsp;APA</th>\n<th>&nbsp;Mailings&nbsp;</th>\n<th>&nbsp;Apazines&nbsp;</th>\n<th>&nbsp;Pages&nbsp;</th</tr>\n"

    allAPAs.sort()
    for apa in allAPAs:
        listText+=(f"\n<tr><td>&nbsp;&nbsp;&nbsp;{FormatLink(APADirName(apa.Name)+'/index.html', apa.Name)}</td>\n"
                          f"<td style='text-align: right'>{apa.Count.Mailings}&nbsp;&nbsp;&nbsp;</td>\n"
                          f"<td style='text-align: right'>{apa.Count.Issues}&nbsp;&nbsp;&nbsp;</td>\n"
                          f"<td style='text-align: right'>{FormatCount(apa.Count.Pages)}&nbsp;&nbsp;&nbsp;</td>\n"
                          f"</tr>\n")
    # Add counts of mailings and contributions to bottom
    for apa in allAPAs:
        allAPAs.Count+=apa.Count
    listText+=(f"\n<tr><td>&nbsp;&nbsp;&nbsp;&nbsp</td>\n"
               f"<td style='text-align: right'>______&nbsp;&nbsp;</td>\n"
               f"<td style='text-align: right'>______&nbsp;&nbsp;</td>\n"
               f"<td style='text-align: right'>______&nbsp;&nbsp;</td>\n")
    listText+=(f"\n<tr><td>&nbsp;&nbsp;&nbsp;&nbsp</td>\n"
               f"<td style='text-align: right'>{allAPAs.Count.Mailings}&nbsp;&nbsp;&nbsp;</td>\n"
               f"<td style='text-align: right'>{allAPAs.Count.Issues}&nbsp;&nbsp;&nbsp;</td>\n"
               f"<td style='text-align: right'>{FormatCount(allAPAs.Count.Pages)}&nbsp;&nbsp;&nbsp;</td>\n"
               f"</tr>\n")

    listText+="</table>\n"
    templateAllApas, success=FindAndReplaceBracketedText(templateAllApas, "fanac-list", listText)

    with open(os.path.join(apaReportsDir, "index.html"), "w", encoding="utf-8") as file:
        file.writelines(templateAllApas)

# End GenerateMailingsReports
###################################################################


# Read the APA Mailings.xlsx file supplied by Joe to get OE, date, etc., information for each mailing.
# The caller has already checked that xlsxname exists.
def ReadXLSX(xlsxname: str, apaName: str) -> dict[str, MailingInfoFromJoe]|None:
    # Read the apa mailings file
    try:
        wb=openpyxl.load_workbook(filename=xlsxname)
    except Exception as e:
        LogError(f"***APA mailings: could not read {xlsxname} ({type(e).__name__}: {e})."
                 f"  The mailing pages will be generated without dates or OEs.")
        return None


    if apaName not in wb.sheetnames:
        return None
    ws=wb[apaName]

    # Separate out the header row
    mailingsheaders=[x.value for x in ws[1]]

    monthCol=FindIndexOfStringInList(mailingsheaders, "Month")
    if monthCol is None:
        LogError(f"{xlsxname} does not contain a 'Month' column")
        return None
    yearCol=FindIndexOfStringInList(mailingsheaders, "Year")
    if yearCol is None:
        LogError(f"{xlsxname} does not contain a 'Year' column")
        return None
    editorCol=FindIndexOfStringInList(mailingsheaders, ["Editor", "OE"])
    if editorCol is None:
        LogError(f"{xlsxname} does not contain an 'Editor' or an 'OE' column")
        return None
    mailingCol=FindIndexOfStringInList(mailingsheaders, ["Mailing", "Issue"])
    if mailingCol is None:
        LogError(f"{xlsxname} does not contain a 'Mailing' or an 'Issue' column")
        return None

    mailingsInfoFromJoe={}
    for i in range(2, 10000):
        row=[x.value for x in ws[i]]
        if all([x is None for x in row]):
            break
        mailingNum=row[mailingCol]
        if type(mailingNum) is int:
            mailingNum=str(mailingNum)  # Standard is to treat mailing number as a string, because sometimes it has to be
        editor=row[editorCol]
        if editor is None:
            editor=""
        mailingsInfoFromJoe[mailingNum]=MailingInfoFromJoe(Number=mailingNum, Year=row[yearCol], Month=row[monthCol], Editor=editor)
    return mailingsInfoFromJoe


######################################################################
# A class to count mailings, issues and pages
class Counts:
    def __init__(self, Pages: int|str=0, Issues: int=0, Mailings: int=0):
        self.Mailings=Mailings
        self.Issues=Issues
        if type(Pages) is str:
            Pages=Int0(Pages)
        self.Pages=Pages

    def __hash__(self):
        return self.Mailings.__hash__()+self.Pages.__hash__()+self.Issues.__hash__()

    def __iadd__(self, val:Counts | int):
        self.Add(val)
        return self

    def __str__(self):
        s=""
        if self.Mailings > 0:
            s+=f"{Pluralize(self.Mailings, 'mailing')}, "
        return s+f"{Pluralize(self.Issues, 'issue')}, {Pluralize(self.Pages, 'page')}"

    # Add a Count or a single fanzine
    def __add__(self, val:Counts | int) -> Counts:
        temp=Counts(Pages=self.Pages, Issues=self.Issues, Mailings=self.Mailings)
        temp.Add(val)
        return temp

    def Add(self, val:Counts | int):
        if type(val) is Counts:
            self.Mailings+=val.Mailings
            self.Issues+=val.Issues
            self.Pages+=val.Pages
            return
        if type(val) is int:
            if self.Mailings == 0:
                self.Mailings=1
            if self.Issues == 0:
                self.Issues=1
            self.Pages+=val
            return
        LogError(f"***Counts.Add() was given a {type(val).__name__} ({val}).  Only a Counts or an int can be added,"
                 f" so this one is ignored and the APA mailing totals will be low by whatever it represented.")


######################################################################
# Entry for a specific mailing in a dictionary of mailings for an APA.
class MailingInfoFromJoe:
    def __init__(self, Number: str = "", Year: str = "", Month: str = "", Editor: str = ""):
        self.Number: str=Number
        self.Editor: str=Editor
        self.Prev: str=""
        self.Next: str=""

        fd=FanzineDate()
        if Month != "":
            fd.Month=Month
        if Year != "":
            fd.Year=Year
        self.Date: FanzineDate=fd

    def __hash__(self):
        return self.Number.__hash__()+self.Editor.__hash__()+self.Prev.__hash__()+self.Next.__hash__()+self.Date.__hash__()


    @property
    def Year(self) -> int:
        return self.Date.Year

    @Year.setter
    def Year(self, y: int) -> None:
        self.Date.Year=y

    @property
    def Month(self) -> int:
        return self.Date.MonthNum

    @Month.setter
    def Month(self, m: int) -> None:
        self.Date.Month=m
# --- end class MailingInfoFromJoe ---


######################################################################
# One mailing of one APA: Joe's information about it, plus the list of apazines in it
class OneMailing:
    def __init__(self):
        self._Count: Counts=Counts()      # The totals for all the apazines in the mailing
        self.MIFJ: MailingInfoFromJoe=MailingInfoFromJoe()       # Joe's info on the mailing
        self.ListFIM: list[FanzineIssueInfo]=[]        # A list of all the apazines in the mailing
        self.Number: str=""        # The name of the mailing (usually a number.)

    def append(self, val: FanzineIssueInfo):
        self.ListFIM.append(val)

    def __str__(self) -> str:
        return self.Number
    def __repr__(self) -> str:
        return self.__str__()
    def __len__(self):
        return len(self.ListFIM)
    def __hash__(self):
        h=self.Number.__hash__()+self.Count.__hash__()+self.MIFJ.__hash__()
        for lf in self.ListFIM:
            h+=lf.__hash__()
        return h

    def __iter__(self):
        self._current=0
        return self

    def __next__(self):
        if self._current >= len(self.ListFIM):
            raise StopIteration
        self._current += 1
        return self.ListFIM[self._current-1]

    def sort(self):
        self.ListFIM.sort(key=lambda x: SortTitle(x.IssueName))

    @property
    def Count(self):
        return self._Count
    @Count.setter
    def Count(self, val):
        self._Count=val



@dataclass
class EntireAPA:
    Count: Counts=field(default_factory=lambda: Counts())
    List: list[OneMailing]=field(default_factory=list)
    Name: str=""

    def __hash__(self):
        h=0
        for om in self.List:
            h+=om.__hash__()
        return h+self.Name.__hash__()+self.Count.__hash__()

    def __len__(self) -> int:
        return len(self.List)

    def append(self, val:OneMailing):
        self.List.append(val)
    def __getitem__(self, index: str) -> OneMailing:
        for (i, x) in enumerate(self.List):
            if x.Number == index:
                return x
        new=OneMailing()
        new.Number=index
        self.List.append(new)
        return new

    def nextIndex(self, index: str) -> str|None:
        for (i, x) in enumerate(self.List):
            if x.Number == index:
                if i+1 >= len(self.List):
                    return None
                return self.List[i+1].Number
        return None

    def prevIndex(self, index: str) -> str|None:
        for (i, x) in enumerate(self.List):
            if x.Number == index:
                if i-1 < 0:
                    return None
                return self.List[i-1].Number
        return None

    def __iter__(self):
        self._current=0
        return self

    def __next__(self):
        if self._current >= len(self.List):
            raise StopIteration
        self._current += 1
        return self.List[self._current-1]

    def sort(self):
        self.List.sort(key=lambda x: SortMessyNumber(x.Number))


@dataclass
class AllAPAs:
    Count: Counts=field(default_factory=lambda: Counts())
    List: list[EntireAPA]=field(default_factory=list)

    def append(self, val:EntireAPA):
        self.List.append(val)

    def __getitem__(self, index: str) -> EntireAPA:
        for (i, x) in enumerate(self.List):
            if x.Name == index:
                return x
        new=EntireAPA()
        new.Name=index
        self.List.append(new)
        return new

    def __iter__(self):
        self._current=0
        return self

    def __next__(self):
        if self._current >= len(self.List):
            raise StopIteration

        self._current += 1
        return self.List[self._current-1]

    def sort(self):
        self.List.sort(key=lambda x: x.Name)
